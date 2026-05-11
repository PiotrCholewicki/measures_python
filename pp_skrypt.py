import openpyxl as xl
import random
from copy import copy
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Font
import os
from openpyxl import Workbook

path1 = 'C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\szablon_pp.xlsx' #ZMIEN LOKALIZACJE
path2 = 'C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\a29.xlsx' #ZMIEN LOKALIZACJE

sheet_name = "a29"  # ZMIEN NA WŁAŚCIWĄ NAZWĘ ARKUSZA

wb1 = xl.load_workbook(filename=path1)
ws1 = wb1.worksheets[0]


page_index = 0
HEADER_INDEX = 12
FOOTER_INDEX = 58
inner_index = HEADER_INDEX + 1 
global_index = inner_index

def add_header(header_text):
    global global_index, inner_index, page_index, ws2
    if(inner_index >= 56):
        inner_index = HEADER_INDEX + 1
        page_index += 1
        copy_rows(ws1, ws2, page_index * 60)
    
    global_index = page_index * 60 + inner_index
    ws2.cell(row=global_index, column=1, value=header_text).font = Font(bold=True, size=16)
    ws2.cell(row=global_index, column=1).alignment = xl.styles.Alignment(horizontal='center')
    ws2.merge_cells(start_row=global_index, start_column=1, end_row=global_index, end_column=9)
    inner_index += 1

def copy_rows(source_sheet, target_sheet, offset):

    for row in source_sheet.iter_rows(min_row=1, max_row=source_sheet.max_row):
        for cell in row:
            target_cell = target_sheet.cell(row=cell.row + offset, column=cell.column, value=cell.value)
            target_cell.font = copy(cell.font)
            target_cell.border = copy(cell.border)
            target_cell.fill = copy(cell.fill)
            target_cell.number_format = copy(cell.number_format)
            target_cell.protection = copy(cell.protection)
            target_cell.alignment = copy(cell.alignment)
            

    ws = target_sheet
    row_number = source_sheet.max_row + offset  # the row that you want to insert page break
    page_break = Break(id=row_number)  # create Break obj
    ws.row_breaks.append(page_break)  # insert page break
    ws.col_breaks.append

def mainloop():
    global global_index, inner_index, page_index, ws2
    while True:
        room_name = input("Podaj nazwę pokoju/STOP: ").upper()
        if(room_name == "SAVE" or room_name == "STOP"):
            wb2.save(path2)
            print("Zapisano plik")
            global_index = page_index * 60 + inner_index
            with open('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_index.txt', 'w') as file:
                file.write(str(global_index))
            break
        else:
            
            # Check if the template isn't fullfilled
            if(inner_index >= 56):
                inner_index = HEADER_INDEX + 1
                page_index += 1
                global_index = page_index * 60 + inner_index
                copy_rows(ws1, ws2, page_index * 60)

            #first add the room name
            global_index = page_index * 60 + inner_index
            if(room_name[0].isdigit()):
                room_name = "POKÓJ " + room_name
            ws2.cell(row = global_index, column = 1, value = room_name).font = Font(bold=True, size=16)
            if(room_name[0:2] == "P "):
                #start new floor on new page
                if(inner_index > 16):
                    inner_index = 57
                add_header("PIĘTRO " + room_name[2:])
                print("Dodano piętro: ", room_name[2:])
                continue
                
            ws2.cell(row = global_index, column = 1).alignment = xl.styles.Alignment(horizontal='center')
            ws2.merge_cells(start_row=global_index, start_column=1, end_row=global_index, end_column=9)
            inner_index += 1

            #now add the sockets with validation
            sockets = input("GNIAZDKA: ")
            if(sockets == "" or sockets.isdigit() == False or int(sockets) < 0):
                if(sockets==""):
                    sockets = 0
                else:
                    while(not sockets.isdigit() or int(sockets) < 0):
                        print("Błędna liczba gniazdek, spróbuj ponownie.")
                        sockets = input("GNIAZDKA: ")

            
            #now add the pc sockets with validation
            pc_sockets = input("GNIAZDKA PC: ")
            if(pc_sockets == "" or pc_sockets.isdigit() == False or int(pc_sockets) < 0):
                if(pc_sockets==""):
                    pc_sockets = 0
                else:
                    while(not pc_sockets.isdigit() or int(pc_sockets) < 0):
                        print("Błędna liczba gniazdek PC, spróbuj ponownie.")
                        pc_sockets = input("GNIAZDKA PC: ")

            sockets_index = 1
            sockets = int(sockets)
            pc_sockets = int(pc_sockets)
            for i in range(sockets):
                #check if the template isn't fullfilled
                if(inner_index >= 57):
                    inner_index = HEADER_INDEX + 1
                    page_index += 1
                    copy_rows(ws1, ws2, page_index * 60)
                    
                #calculate the global index1
                global_index = page_index * 60 + inner_index
                
                ws2.cell(row=global_index, column=1, value=sockets_index)
                ws2.cell(row=global_index, column=2, value="Gniazdko 1-fazowe")
                ws2.cell(row=global_index, column=3, value="SBinst")
                ws2.cell(row=global_index, column=4, value="B16")
                ws2.cell(row=global_index, column=5, value=230.00)
                ws2.cell(row=global_index, column=6, value="=RAND()*0.35+0.9")
                ws2.cell(row=global_index, column=7, value=80)
                ws2.cell(row=global_index, column=8, value=f"=F{global_index}*G{global_index}")  
                ws2.cell(row=global_index, column=9, value="TAK")
                sockets_index += 1
                inner_index += 1
                
                
            for i in range(pc_sockets):
                #check if the template isn't fullfilled
                if(inner_index >= 57):
                    page_index += 1
                    inner_index = HEADER_INDEX + 1
                    copy_rows(ws1, ws2, page_index * 60)
                    
                global_index = page_index * 60 + inner_index
                ws2.cell(row=global_index, column=1, value=sockets_index)
                ws2.cell(row=global_index, column=2, value="Gniazdko 1-fazowe PC")
                ws2.cell(row=global_index, column=3, value="SBinst")
                ws2.cell(row=global_index, column=4, value="B16")
                ws2.cell(row=global_index, column=5, value=230.00)
                ws2.cell(row=global_index, column=6, value="=RAND()*0.35+0.9")
                ws2.cell(row=global_index, column=7, value=80)
                ws2.cell(row=global_index, column=8, value=f"=F{global_index}*G{global_index}")  
                ws2.cell(row=global_index, column=9, value="TAK")
                sockets_index += 1
                inner_index += 1
                
        wb2.save(path2)
        global_index = page_index * 60 + inner_index
        sockets_index = 1
        with open('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_index.txt', 'w') as file:
            file.write(str(global_index))
        with open('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_room.txt', 'w') as file:
            file.write(room_name)
            
            


if not os.path.exists(path2):
    wb = Workbook()
    wb.save(path2)
    print("Utworzono plik izo:", path2)

if not os.path.exists('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_index.txt'):
    print("No last index file found, starting fresh.")
    
    wb2 = xl.load_workbook(filename=path2)
    ws2 = wb2.create_sheet(sheet_name)
    #initial copy of a template, if it's the first run
    ws2.column_dimensions['A'].width = 8.43
    ws2.column_dimensions['B'].width = 40.43
    ws2.column_dimensions['C'].width = 20.43
    ws2.column_dimensions['D'].width = 9.57
    ws2.column_dimensions['E'].width = 23.29
    ws2.column_dimensions['F'].width = 16.14
    ws2.column_dimensions['G'].width = 16.86
    ws2.column_dimensions['H'].width = 8.8
    ws2.column_dimensions['I'].width = 13.14
    copy_rows(ws1, ws2, 0)
    mainloop()

else:
    print(f"The file 'last_index.txt' already exists.")
    with open('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_index.txt', 'r') as file:
        global_index = int(file.read().strip())
        print(f"Last index loaded: {global_index}")
    with open('C:\\Users\\Piotr\\Desktop\\mieszkania\\skrypt_pp\\last_room.txt', 'r') as file:
        last_room = file.read().strip()
        print(f"Last room loaded: {last_room}")
    wb2 = xl.load_workbook(filename=path2)
    inner_index = global_index % 60
    page_index = global_index // 60
    if sheet_name in wb2.sheetnames:
        ws2 = wb2[sheet_name]
        mainloop()
    else:
        ws2 = wb2.create_sheet(sheet_name)




